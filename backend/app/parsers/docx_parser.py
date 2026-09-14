"""DOCX document parser using python-docx."""

import io
from pathlib import Path
from typing import List
import docx
from backend.app.parsers.base_parser import BaseParser, ParsedPage


class DOCXParser(BaseParser):
    """Parses DOCX documents into paragraphs/tables."""

    def parse_file(self, file_path: Path) -> List[ParsedPage]:
        with open(file_path, "rb") as f:
            return self.parse_bytes(f.read(), file_path.name)

    def parse_bytes(self, content: bytes, filename: str) -> List[ParsedPage]:
        stream = io.BytesIO(content)
        doc = docx.Document(stream)

        text_parts: List[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n\n".join(text_parts)
        return [ParsedPage(page_number=1, text=full_text)]
